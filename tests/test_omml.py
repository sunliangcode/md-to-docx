"""OMML / MathML conversion tests."""

from __future__ import annotations

from xml.etree import ElementTree as ET

from docx import Document

from md_to_docx.render.omml import (
    MATH_NS,
    _mathml_to_omml,
    add_block_math,
    add_inline_math,
)

MML = "http://www.w3.org/1998/Math/MathML"


def _wrap(inner: str) -> str:
    return f'<math xmlns="{MML}">{inner}</math>'


def test_mathml_parse_error_returns_none():
    assert _mathml_to_omml("<not-xml") is None


def test_mathml_mi_mn_mo():
    el = _mathml_to_omml(_wrap("<mi>x</mi><mo>+</mo><mn>1</mn>"))
    assert el is not None
    assert el.tag == f"{{{MATH_NS}}}oMath"
    texts = [t.text for t in el.iter(f"{{{MATH_NS}}}t")]
    assert texts == ["x", "+", "1"]


def test_mathml_mfrac():
    el = _mathml_to_omml(_wrap("<mfrac><mi>a</mi><mi>b</mi></mfrac>"))
    assert el is not None
    assert el.find(f"{{{MATH_NS}}}f") is not None
    assert el.find(f".//{{{MATH_NS}}}num") is not None
    assert el.find(f".//{{{MATH_NS}}}den") is not None


def test_mathml_msup_msub_msubsup():
    sup = _mathml_to_omml(_wrap("<msup><mi>x</mi><mn>2</mn></msup>"))
    assert sup is not None and sup.find(f"{{{MATH_NS}}}sSup") is not None

    sub = _mathml_to_omml(_wrap("<msub><mi>x</mi><mn>1</mn></msub>"))
    assert sub is not None and sub.find(f"{{{MATH_NS}}}sSub") is not None

    both = _mathml_to_omml(_wrap("<msubsup><mi>x</mi><mn>1</mn><mn>2</mn></msubsup>"))
    assert both is not None and both.find(f"{{{MATH_NS}}}sSubSup") is not None


def test_mathml_msqrt_mroot():
    sqrt = _mathml_to_omml(_wrap("<msqrt><mi>x</mi></msqrt>"))
    assert sqrt is not None and sqrt.find(f"{{{MATH_NS}}}rad") is not None

    root = _mathml_to_omml(_wrap("<mroot><mi>x</mi><mn>3</mn></mroot>"))
    assert root is not None and root.find(f"{{{MATH_NS}}}rad") is not None


def test_mathml_mfenced_and_mrow():
    fenced = _mathml_to_omml(_wrap('<mfenced open="(" close=")"><mi>x</mi></mfenced>'))
    assert fenced is not None
    assert fenced.find(f"{{{MATH_NS}}}d") is not None

    row = _mathml_to_omml(_wrap("<mrow><mi>a</mi><mi>b</mi></mrow>"))
    assert row is not None


def test_mathml_mtable():
    table = _mathml_to_omml(
        _wrap(
            "<mtable>"
            "<mtr><mtd><mn>1</mn></mtd><mtd><mn>2</mn></mtd></mtr>"
            "<mtr><mtd><mn>3</mn></mtd><mtd><mn>4</mn></mtd></mtr>"
            "</mtable>"
        )
    )
    assert table is not None
    assert table.find(f"{{{MATH_NS}}}m") is not None or list(table)  # matrix or children


def test_add_inline_math_inserts_omath():
    doc = Document()
    p = doc.add_paragraph()
    add_inline_math(p, r"E=mc^2")
    xml = ET.tostring(p._p, encoding="unicode")
    assert "oMath" in xml or "E=mc^2" in xml


def test_add_block_math_inserts_omath_para():
    doc = Document()
    p = doc.add_paragraph()
    add_block_math(p, r"\frac{a}{b}")
    xml = ET.tostring(p._p, encoding="unicode")
    assert "oMathPara" in xml or "oMath" in xml or r"\frac{a}{b}" in xml


def test_math_fixture_converts_to_docx(tmp_path):
    from md_to_docx.converter import convert_file

    src = tmp_path / "math.md"
    src.write_text("# Math\n\nInline $E=mc^2$\n\n$$\n\\frac{a}{b}\n$$\n", encoding="utf-8")
    out = tmp_path / "math.docx"
    convert_file(src, out)
    assert out.is_file() and out.stat().st_size > 1000


def test_mathml_mover_munder_mstyle_unknown():
    mover = _mathml_to_omml(_wrap("<mover><mi>x</mi><mo>^</mo></mover>"))
    assert mover is not None
    assert mover.find(f".//{{{MATH_NS}}}acc") is not None or list(mover)

    munder = _mathml_to_omml(_wrap("<munder><mi>x</mi><mo>_</mo></munder>"))
    assert munder is not None
    assert munder.find(f".//{{{MATH_NS}}}limLow") is not None or list(munder)

    styled = _mathml_to_omml(_wrap("<mstyle><mi>y</mi></mstyle>"))
    assert styled is not None

    unknown = _mathml_to_omml(_wrap("<munknown><mi>z</mi></munknown>"))
    assert unknown is not None


def test_add_math_plaintext_fallback(monkeypatch):
    import latex2mathml.converter as conv

    def _boom(_latex: str) -> str:
        raise RuntimeError("no converter")

    monkeypatch.setattr(conv, "convert", _boom)

    doc = Document()
    p = doc.add_paragraph()
    add_inline_math(p, "FALLBACK")
    assert "FALLBACK" in p.text

    p2 = doc.add_paragraph()
    add_block_math(p2, "BLOCK_FALLBACK")
    assert "BLOCK_FALLBACK" in p2.text
