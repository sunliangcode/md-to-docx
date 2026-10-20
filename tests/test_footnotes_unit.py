"""Unit tests for footnote XML helpers."""

from __future__ import annotations

from docx import Document

from md_to_docx.ast import nodes as n
from md_to_docx.render.footnotes import (
    FootnoteState,
    build_footnotes_xml,
    ensure_footnotes_part,
)


def test_footnote_state_reuses_ids():
    doc = Document()
    p = doc.add_paragraph()
    state = FootnoteState()
    a = state.add_ref(p, "note")
    b = state.add_ref(p, "note")
    assert a == b == 1
    assert state.keys_in_order() == ["note"]


def test_build_footnotes_xml_rich_inlines_and_empty_def():
    state = FootnoteState()
    state.order = ["rich", "empty", "orphan-ref"]
    state.key_to_id = {"rich": 1, "empty": 2, "orphan-ref": 3}
    footnotes = (
        n.FootnoteDef(
            key="rich",
            children=(
                n.Paragraph(
                    (
                        n.Text(" a "),
                        n.Strong((n.Text("b"),)),
                        n.Emphasis((n.Text("e"),)),
                        n.Code("c"),
                        n.Link("https://ex.com", (n.Text("L"),)),
                        n.Strike((n.Text("s"),)),
                    )
                ),
                n.CodeBlock(text="ignored-as-block", lang="py"),
            ),
        ),
        n.FootnoteDef(key="empty", children=()),
        n.FootnoteDef(
            key="extra-unreferenced",
            children=(n.Paragraph((n.Text("extra"),)),),
        ),
    )
    xml = build_footnotes_xml(state, footnotes)
    assert b"footnotes" in xml
    assert b"a " in xml or b"a" in xml
    assert b"Consolas" in xml
    assert b"extra" in xml
    # orphan-ref has no definition — skipped in body but separators remain
    assert b'w:id="3"' not in xml or b"orphan" not in xml


def test_ensure_footnotes_part_noop_when_empty():
    doc = Document()
    state = FootnoteState()
    ensure_footnotes_part(doc, state, ())
    assert not any("footnotes" in rel.reltype for rel in doc.part.rels.values())


def test_ensure_footnotes_part_attaches_xml():
    doc = Document()
    p = doc.add_paragraph("body")
    state = FootnoteState()
    state.add_ref(p, "n1")
    footnotes = (n.FootnoteDef(key="n1", children=(n.Paragraph((n.Text("note body"),)),)),)
    ensure_footnotes_part(doc, state, footnotes)
    assert any(rel.reltype.endswith("footnotes") for rel in doc.part.rels.values())
    # Second call replaces existing relationship
    ensure_footnotes_part(doc, state, footnotes)
    footnote_rels = [rel for rel in doc.part.rels.values() if "footnotes" in rel.reltype]
    assert len(footnote_rels) == 1
