"""Native DOCX style configuration tests."""

from __future__ import annotations

from docx import Document
from docx.shared import RGBColor

from md_to_docx.render.styles import configure_document_styles, set_theme_font_lang


def test_configure_document_styles_creates_custom_styles():
    doc = Document()
    configure_document_styles(doc, force=True, heading_color=RGBColor(0x11, 0x18, 0x27))
    assert "MDCodeBlock" in [s.name for s in doc.styles]
    assert "Caption" in [s.name for s in doc.styles]
    assert doc.styles["Heading 1"].font.color.rgb == RGBColor(0x11, 0x18, 0x27)


def test_configure_document_styles_idempotent_without_force():
    doc = Document()
    configure_document_styles(doc, force=True)
    configure_document_styles(doc, force=False)
    assert doc.styles["MDCodeBlock"].font.name == "Consolas"


def test_set_theme_font_lang_twice():
    doc = Document()
    set_theme_font_lang(doc, latin="en-US", east_asia="zh-CN")
    set_theme_font_lang(doc, latin="en-GB", east_asia="ja-JP")
    settings = doc.settings.element
    from docx.oxml.ns import qn

    theme = settings.find(qn("w:themeFontLang"))
    assert theme is not None
    assert theme.get(qn("w:val")) == "en-GB"
    assert theme.get(qn("w:eastAsia")) == "ja-JP"
