"""Image embedding helpers."""

from __future__ import annotations

from pathlib import Path

from docx.document import Document as DocxDocument
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.shared import Inches


def embed_image_path(
    doc: DocxDocument, path: Path, *, alt: str = "", width_inches: float = 5.0
) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    run.add_picture(str(path), width=Inches(width_inches))
    if alt:
        p.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
